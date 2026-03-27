from __future__ import annotations

import io
import json
import logging
from datetime import date as current_date
from pathlib import Path
from typing import Any
from urllib.parse import quote

from docx import Document
from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse

from .asr import ASRServiceError, ASRTimeoutError, transcribe_audio_with_speakers
from .graph import meeting_workflow, render_report_markdown
from .history_store import (
    clear_meeting_history,
    delete_action_item,
    delete_meeting_history,
    get_meeting_history_record,
    init_history_db,
    list_action_items,
    list_meeting_history,
    replace_action_items_for_history,
    save_meeting_history,
    update_action_item_status,
    update_audio_transcript_history,
)
from .schemas import (
    ActionItem,
    ActionItemRecord,
    ActionItemStatusUpdateRequest,
    AnalyzeResult,
    AnalyzeTextRequest,
    AudioTranscriptDownloadRequest,
    AudioTranscriptResponse,
    MeetingHistoryItem,
    WeeklyReportRequest,
    WeeklyReportResponse,
)

load_dotenv()

logger = logging.getLogger(__name__)
if not logger.handlers:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")

app = FastAPI(title="Meeting Agent API", version="0.9.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

AUDIO_EXTENSIONS = {"mp3", "wav", "m4a", "mp4", "aac", "ogg", "webm"}
ENTRY_TYPE_MEETING_ANALYSIS = "meeting_analysis"
ENTRY_TYPE_AUDIO_TRANSCRIPT = "audio_transcript"
MEETING_DOWNLOAD_OPTIONS = {
    "management-markdown",
    "management-docx",
    "project-markdown",
    "project-docx",
}
AUDIO_DOWNLOAD_OPTIONS = {"transcript-markdown", "transcript-docx"}
AUDIO_MINUTES_REPORT_TYPE = "project"


@app.on_event("startup")
def startup() -> None:
    init_history_db()


def _safe_title(value: str | None, fallback: str) -> str:
    normalized = str(value or "").strip()
    return normalized or fallback


def _build_source_title(filename: str | None, fallback: str) -> str:
    raw_name = Path(filename or "").stem.strip()
    return raw_name or fallback


def _build_download_headers(filename: str) -> dict[str, str]:
    ascii_filename = "".join(ch if ch.isascii() and ch not in '\\/:*?"<>|' else "_" for ch in filename)
    ascii_filename = ascii_filename.strip(" ._") or "report"
    return {
        "Content-Disposition": f"attachment; filename={ascii_filename}; filename*=UTF-8''{quote(filename)}",
    }


def _build_report_state(request: WeeklyReportRequest) -> dict[str, Any]:
    return {
        "raw_text": "",
        "title": request.title,
        "date": request.date or "",
        "weekly_period": request.weekly_period,
        "decisions": [item for item in request.decisions],
        "actions": [item.model_dump() for item in request.actions],
        "risks": [item for item in request.risks],
        "report_type": request.report_type,
        "validation": {
            "passed": True,
            "estimated_accuracy": 1.0,
            "estimated_action_recall": 1.0,
            "missing_action_hints": [],
            "ambiguous_owners": [],
            "notes": [],
            "iteration": 0,
        },
    }


def _render_transcript_markdown(title: str, transcript_text: str) -> str:
    return f"# {title}\n\n## 转写全文\n\n{transcript_text.strip()}\n"


def _build_transcript_docx_bytes(request: AudioTranscriptDownloadRequest) -> io.BytesIO:
    document = Document()
    document.add_heading(request.title, level=1)
    document.add_heading("转写全文", level=2)

    paragraphs = [item.strip() for item in request.transcript_text.splitlines() if item.strip()]
    if paragraphs:
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    else:
        document.add_paragraph(request.transcript_text.strip())

    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output


def _render_text_meeting_markdown(request: WeeklyReportRequest) -> str:
    report_markdown = render_report_markdown(_build_report_state(request)).strip()
    source_text = str(getattr(request, "source_text", "") or "").strip()
    if not source_text:
        return report_markdown + "\n"

    return "\n".join(
        [
            "## \u539f\u59cb\u4f1a\u8bae\u5185\u5bb9",
            "",
            source_text,
            "",
            "## \u751f\u6210\u7eaa\u8981",
            "",
            report_markdown,
        ]
    ).strip() + "\n"


def _build_weekly_docx_bytes(request: WeeklyReportRequest) -> io.BytesIO:
    document = Document()
    document_title = "\u9879\u76ee\u7ba1\u7406\u5468\u62a5" if request.report_type == "management" else _safe_title(request.title, "\u9879\u76ee\u4f1a\u8bae\u7eaa\u8981")
    document.add_heading(document_title, level=1)

    if request.date:
        document.add_paragraph(f"\u65e5\u671f\uff1a{request.date}")
    if request.weekly_period:
        document.add_paragraph(request.weekly_period)

    source_text = str(getattr(request, "source_text", "") or "").strip()
    content_heading_level = 3 if source_text else 2

    if source_text:
        document.add_heading("\u539f\u59cb\u4f1a\u8bae\u5185\u5bb9", level=2)
        source_paragraphs = [item.strip() for item in source_text.splitlines() if item.strip()]
        if source_paragraphs:
            for paragraph in source_paragraphs:
                document.add_paragraph(paragraph)
        else:
            document.add_paragraph(source_text)
        document.add_heading("\u751f\u6210\u7eaa\u8981", level=2)

    document.add_heading("\u51b3\u7b56\u9879", level=content_heading_level)
    if request.decisions:
        for decision in request.decisions:
            document.add_paragraph(decision, style="List Bullet")
    else:
        document.add_paragraph("\u6682\u65e0")

    document.add_heading("\u884c\u52a8\u9879", level=content_heading_level)
    if request.actions:
        for action in request.actions:
            document.add_paragraph(
                f"{action.task} | \u8d1f\u8d23\u4eba\uff1a{action.owner} | \u622a\u6b62\u65e5\u671f\uff1a{action.deadline}",
                style="List Bullet",
            )
    else:
        document.add_paragraph("\u6682\u65e0")

    document.add_heading("\u98ce\u9669\u9879", level=content_heading_level)
    if request.risks:
        for risk in request.risks:
            document.add_paragraph(risk, style="List Bullet")
    else:
        document.add_paragraph("\u6682\u65e0")

    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output


def _build_weekly_request_from_payload(
    payload: dict[str, Any], report_type: str, title_override: str | None = None
) -> WeeklyReportRequest:
    actions = [ActionItem(**action) for action in payload.get("actions", [])]
    return WeeklyReportRequest(
        title=_safe_title(title_override or payload.get("title"), "\u4f1a\u8bae\u7eaa\u8981"),
        date=str(payload.get("date", "") or ""),
        weekly_period=str(payload.get("weekly_period", "") or ""),
        decisions=[str(item) for item in payload.get("decisions", [])],
        actions=actions,
        risks=[str(item) for item in payload.get("risks", [])],
        report_type=report_type,
        source_text=str(payload.get("source_text", "") or ""),
    )


def _build_audio_minutes_request(
    title: str,
    transcript_text: str,
    analysis_payload: dict[str, Any] | None = None,
    report_type: str = AUDIO_MINUTES_REPORT_TYPE,
) -> AudioTranscriptDownloadRequest:
    payload = analysis_payload or {}
    actions = [
        action if isinstance(action, ActionItem) else ActionItem(**action)
        for action in payload.get("actions", [])
    ]
    return AudioTranscriptDownloadRequest(
        title=_safe_title(title, "会议语音纪要"),
        transcript_text=transcript_text,
        date=str(payload.get("date", "") or ""),
        weekly_period=str(payload.get("weekly_period", "") or ""),
        decisions=[str(item) for item in payload.get("decisions", [])],
        actions=actions,
        risks=[str(item) for item in payload.get("risks", [])],
        report_type=report_type,
    )



def _render_audio_minutes_markdown(request: AudioTranscriptDownloadRequest) -> str:
    lines = [
        f"# {request.title}",
        "",
        "## 语音转写全文",
        "",
        request.transcript_text.strip() or "暂无",
        "",
        "## 提取纪要",
    ]

    if request.date:
        lines.extend(["", f"日期：{request.date}"])
    if request.weekly_period:
        lines.extend(["", request.weekly_period])

    lines.extend(["", "### 决策项"])
    if request.decisions:
        lines.extend([f"- {item}" for item in request.decisions])
    else:
        lines.append("- 暂无")

    lines.extend(["", "### 行动项"])
    if request.actions:
        for action in request.actions:
            risk_text = str(action.risk or "").strip()
            suffix = ""
            if risk_text and risk_text.lower() not in {"none", "无"}:
                suffix = f" | 风险：{risk_text}"
            lines.append(
                f"- {action.task} | 负责人：{action.owner} | 截止日期：{action.deadline}{suffix}"
            )
    else:
        lines.append("- 暂无")

    lines.extend(["", "### 风险项"])
    if request.risks:
        lines.extend([f"- {item}" for item in request.risks])
    else:
        lines.append("- 暂无")

    return "\n".join(lines).strip() + "\n"



def _build_audio_minutes_docx_bytes(request: AudioTranscriptDownloadRequest) -> io.BytesIO:
    document = Document()
    document.add_heading(request.title, level=1)

    document.add_heading("语音转写全文", level=2)
    paragraphs = [item.strip() for item in request.transcript_text.splitlines() if item.strip()]
    if paragraphs:
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    else:
        document.add_paragraph(request.transcript_text.strip() or "暂无")

    document.add_heading("提取纪要", level=2)
    if request.date:
        document.add_paragraph(f"日期：{request.date}")
    if request.weekly_period:
        document.add_paragraph(request.weekly_period)

    document.add_heading("决策项", level=3)
    if request.decisions:
        for decision in request.decisions:
            document.add_paragraph(decision, style="List Bullet")
    else:
        document.add_paragraph("暂无")

    document.add_heading("行动项", level=3)
    if request.actions:
        for action in request.actions:
            risk_text = str(action.risk or "").strip()
            suffix = ""
            if risk_text and risk_text.lower() not in {"none", "无"}:
                suffix = f" | 风险：{risk_text}"
            document.add_paragraph(
                f"{action.task} | 负责人：{action.owner} | 截止日期：{action.deadline}{suffix}",
                style="List Bullet",
            )
    else:
        document.add_paragraph("暂无")

    document.add_heading("风险项", level=3)
    if request.risks:
        for risk in request.risks:
            document.add_paragraph(risk, style="List Bullet")
    else:
        document.add_paragraph("暂无")

    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output


def _load_structured_data(record: dict[str, Any]) -> dict[str, Any]:
    raw_value = record.get("structured_data") or "{}"
    try:
        payload = json.loads(raw_value)
        return payload if isinstance(payload, dict) else {}
    except json.JSONDecodeError:
        logger.warning("Failed to decode structured_data for history id=%s", record.get("id"))
        return {}


def _normalize_action_payloads(actions: list[Any] | None) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    for action in actions or []:
        if isinstance(action, ActionItem):
            payload = action.model_dump()
        elif isinstance(action, dict):
            payload = dict(action)
        else:
            try:
                payload = ActionItem(**dict(action)).model_dump()
            except Exception:
                continue

        task = str(payload.get("task", "") or "").strip()
        if not task:
            continue

        normalized.append(
            {
                "task": task,
                "owner": str(payload.get("owner", "") or "").strip() or "\u5f85\u5b9a",
                "deadline": str(payload.get("deadline", "") or "").strip() or "\u5f85\u5b9a",
                "risk": str(payload.get("risk", "") or "").strip() or "\u65e0",
                "status": str(payload.get("status", "") or "pending").strip() or "pending",
            }
        )
    return normalized


def _sync_action_items_for_history(
    history_id: int,
    *,
    title: str,
    date: str,
    actions: list[Any] | None,
) -> None:
    replace_action_items_for_history(
        meeting_history_id=history_id,
        title=_safe_title(title, "\u4f1a\u8bae\u7eaa\u8981"),
        date=str(date or ""),
        actions=_normalize_action_payloads(actions),
    )


def _save_meeting_analysis_history(payload: dict[str, Any], report_type: str, title_override: str | None = None) -> int:
    display_title = _safe_title(title_override or payload.get("title"), "会议纪要")
    weekly_request = _build_weekly_request_from_payload(payload, report_type, title_override=display_title)
    markdown_content = _render_text_meeting_markdown(weekly_request)
    docx_content = _build_weekly_docx_bytes(weekly_request).getvalue()
    structured_data = json.dumps(
        {
            "display_title": display_title,
            "generated_title": payload.get("title", ""),
            "date": weekly_request.date,
            "weekly_period": weekly_request.weekly_period,
            "decisions": payload.get("decisions", []),
            "actions": payload.get("actions", []),
            "risks": payload.get("risks", []),
            "validation": payload.get("validation", {}),
            "source_text": weekly_request.source_text,
        },
        ensure_ascii=False,
    )
    history_id = save_meeting_history(
        title=display_title,
        date=weekly_request.date,
        entry_type=ENTRY_TYPE_MEETING_ANALYSIS,
        report_type=report_type,
        structured_data=structured_data,
        markdown_content=markdown_content,
        docx_content=docx_content,
    )
    logger.info(
        "Saved meeting analysis history: id=%s title=%s report_type=%s",
        history_id,
        display_title,
        report_type,
    )
    return history_id


def _save_audio_transcript_history(title: str, transcript_text: str, analysis_payload: dict[str, Any] | None = None) -> int:
    display_title = _safe_title(title, "\u4f1a\u8bae\u8bed\u97f3\u7eaa\u8981")
    audio_request = _build_audio_minutes_request(display_title, transcript_text, analysis_payload)
    markdown_content = _render_audio_minutes_markdown(audio_request)
    docx_content = _build_audio_minutes_docx_bytes(audio_request).getvalue()
    analysis_payload = analysis_payload or {}
    structured_data = json.dumps(
        {
            "display_title": display_title,
            "date": audio_request.date or str(current_date.today()),
            "weekly_period": audio_request.weekly_period,
            "transcript_text": transcript_text,
            "decisions": [item for item in audio_request.decisions],
            "actions": [action.model_dump() for action in audio_request.actions],
            "risks": [item for item in audio_request.risks],
            "validation": analysis_payload.get("validation", {}),
            "report_markdown": analysis_payload.get("report_markdown", ""),
            "report_type": audio_request.report_type,
        },
        ensure_ascii=False,
    )
    history_id = save_meeting_history(
        title=display_title,
        date=audio_request.date or str(current_date.today()),
        entry_type=ENTRY_TYPE_AUDIO_TRANSCRIPT,
        report_type=audio_request.report_type,
        structured_data=structured_data,
        markdown_content=markdown_content,
        docx_content=docx_content,
    )
    logger.info("Saved audio meeting history: id=%s title=%s", history_id, display_title)
    return history_id


def _update_audio_transcript_history_entry(
    history_id: int,
    title: str,
    transcript_text: str,
    analysis_payload: dict[str, Any] | None = None,
) -> bool:
    record = get_meeting_history_record(history_id)
    if not record or str(record.get("entry_type") or "") != ENTRY_TYPE_AUDIO_TRANSCRIPT:
        return False

    display_title = _safe_title(title, "\u4f1a\u8bae\u8bed\u97f3\u7eaa\u8981")
    audio_request = _build_audio_minutes_request(display_title, transcript_text, analysis_payload)
    markdown_content = _render_audio_minutes_markdown(audio_request)
    docx_content = _build_audio_minutes_docx_bytes(audio_request).getvalue()
    structured_data = _load_structured_data(record)
    structured_data.update(
        {
            "display_title": display_title,
            "date": audio_request.date or structured_data.get("date") or str(current_date.today()),
            "weekly_period": audio_request.weekly_period,
            "transcript_text": transcript_text,
            "decisions": [item for item in audio_request.decisions],
            "actions": [action.model_dump() for action in audio_request.actions],
            "risks": [item for item in audio_request.risks],
            "validation": (analysis_payload or {}).get("validation", {}),
            "report_markdown": (analysis_payload or {}).get("report_markdown", ""),
            "report_type": audio_request.report_type,
        }
    )

    updated = update_audio_transcript_history(
        history_id,
        title=display_title,
        structured_data=json.dumps(structured_data, ensure_ascii=False),
        markdown_content=markdown_content,
        docx_content=docx_content,
    )
    if updated:
        logger.info("Updated audio meeting history: id=%s title=%s", history_id, display_title)
    return updated


def _run_workflow(text: str, report_type: str = "management", max_iterations: int = 2) -> dict[str, Any]:
    state: dict[str, Any] = {
        "raw_text": text,
        "report_type": report_type,
        "max_iterations": max_iterations,
        "iteration": 0,
    }
    result = meeting_workflow.invoke(state)
    return {
        "title": result.get("title", "会议结构化提取结果"),
        "date": result.get("date", ""),
        "weekly_period": result.get("weekly_period", ""),
        "decisions": result.get("decisions", []),
        "actions": result.get("actions", []),
        "risks": result.get("risks", []),
        "validation": result.get("validation", {}),
        "report_markdown": result.get("report_markdown", ""),
    }


def _download_bytes_response(content: bytes, media_type: str, filename: str) -> StreamingResponse:
    return StreamingResponse(io.BytesIO(content), media_type=media_type, headers=_build_download_headers(filename))


def _download_meeting_analysis_record(record: dict[str, Any], option: str) -> StreamingResponse:
    if option not in MEETING_DOWNLOAD_OPTIONS:
        raise HTTPException(status_code=400, detail="会议解析记录仅支持管理版/项目版的 Markdown 或 Word 下载")

    report_type, file_format = option.split("-", 1)
    structured_data = _load_structured_data(record)
    weekly_request = _build_weekly_request_from_payload(structured_data, report_type, title_override=str(record.get("title") or "会议纪要"))
    report_label = "管理版" if report_type == "management" else "项目版"

    if file_format == "markdown":
        markdown = _render_text_meeting_markdown(weekly_request)
        filename = f"{weekly_request.title}_{report_label}.md"
        return _download_bytes_response(markdown.encode("utf-8"), "text/markdown; charset=utf-8", filename)

    docx_bytes = _build_weekly_docx_bytes(weekly_request).getvalue()
    filename = f"{weekly_request.title}_{report_label}.docx"
    return _download_bytes_response(
        docx_bytes,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename,
    )



def _download_audio_transcript_record(record: dict[str, Any], option: str) -> StreamingResponse:
    if option not in AUDIO_DOWNLOAD_OPTIONS:
        raise HTTPException(status_code=400, detail="\u8bed\u97f3\u7eaa\u8981\u8bb0\u5f55\u4ec5\u652f\u6301 Markdown \u6216 Word \u4e0b\u8f7d")

    title = _safe_title(record.get("title"), "\u4f1a\u8bae\u8bed\u97f3\u7eaa\u8981")
    if option == "transcript-markdown":
        filename = f"{title}_\u8bed\u97f3\u7eaa\u8981.md"
        content = str(record.get("markdown_content") or "")
        return _download_bytes_response(content.encode("utf-8"), "text/markdown; charset=utf-8", filename)

    filename = f"{title}_\u8bed\u97f3\u7eaa\u8981.docx"
    docx_content = bytes(record.get("docx_content") or b"")
    return _download_bytes_response(
        docx_content,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename,
    )


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/api/analyze/text", response_model=AnalyzeResult)
def analyze_text(request: AnalyzeTextRequest) -> AnalyzeResult:
    payload = _run_workflow(request.text, report_type=request.report_type)
    payload["source_text"] = request.text
    result = AnalyzeResult(**payload)
    history_id = _save_meeting_analysis_history(result.model_dump(), request.report_type)
    _sync_action_items_for_history(
        history_id,
        title=_safe_title(result.title, "\u4f1a\u8bae\u7eaa\u8981"),
        date=result.date,
        actions=result.actions,
    )
    return result


@app.post("/api/analyze/upload", response_model=AnalyzeResult)
async def analyze_upload(file: UploadFile = File(...), report_type: str = "management") -> AnalyzeResult:
    filename = file.filename or ""
    ext = filename.lower().split(".")[-1] if "." in filename else ""
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="上传文件为空")

    if ext in {"txt", "md"}:
        text = content.decode("utf-8", errors="ignore")
    elif ext == "docx":
        document = Document(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    else:
        raise HTTPException(status_code=400, detail="仅支持 txt、md、docx 文件")

    payload = _run_workflow(text, report_type=report_type)
    payload["source_text"] = text
    result = AnalyzeResult(**payload)
    history_title = _build_source_title(filename, _safe_title(result.title, "会议纪要"))
    history_id = _save_meeting_analysis_history(result.model_dump(), report_type, title_override=history_title)
    _sync_action_items_for_history(
        history_id,
        title=history_title,
        date=result.date,
        actions=result.actions,
    )
    return result


@app.post("/api/audio/transcribe", response_model=AudioTranscriptResponse)
async def transcribe_audio(file: UploadFile = File(...)) -> AudioTranscriptResponse:
    filename = file.filename or "meeting_audio.wav"
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext not in AUDIO_EXTENSIONS:
        raise HTTPException(status_code=400, detail="\u4ec5\u652f\u6301 mp3\u3001wav\u3001m4a\u3001mp4\u3001aac\u3001ogg\u3001webm \u97f3\u9891\u6587\u4ef6")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="\u4e0a\u4f20\u97f3\u9891\u6587\u4ef6\u4e3a\u7a7a")

    try:
        transcript_text = transcribe_audio_with_speakers(filename, content)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except ASRTimeoutError as exc:
        raise HTTPException(status_code=504, detail=str(exc)) from exc
    except ASRServiceError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    except Exception as exc:
        logger.exception("Unexpected ASR failure: filename=%s", filename)
        raise HTTPException(status_code=502, detail=f"\u8bed\u97f3\u8f6c\u5199\u670d\u52a1\u5f02\u5e38\uff1a{exc}") from exc

    title = _build_source_title(filename, "\u4f1a\u8bae\u8bed\u97f3\u7eaa\u8981")
    return AudioTranscriptResponse(
        title=title,
        file_name=filename,
        transcript_text=transcript_text,
        markdown=_render_transcript_markdown(title, transcript_text),
        history_id=None,
        report_type=AUDIO_MINUTES_REPORT_TYPE,
    )


@app.post("/api/audio/transcribe/markdown")
def download_audio_transcript_markdown(request: AudioTranscriptDownloadRequest) -> StreamingResponse:
    markdown = _render_audio_minutes_markdown(request)
    filename = f"{request.title}_\u8bed\u97f3\u7eaa\u8981.md"
    return _download_bytes_response(markdown.encode("utf-8"), "text/markdown; charset=utf-8", filename)


@app.post("/api/audio/transcribe/docx")
def download_audio_transcript_docx(request: AudioTranscriptDownloadRequest) -> StreamingResponse:
    filename = f"{request.title}_\u8bed\u97f3\u7eaa\u8981.docx"
    return _download_bytes_response(
        _build_audio_minutes_docx_bytes(request).getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename,
    )


@app.post("/api/audio/finalize", response_model=AudioTranscriptResponse)
def finalize_audio_minutes(request: AudioTranscriptDownloadRequest) -> AudioTranscriptResponse:
    analysis_payload = _run_workflow(request.transcript_text, report_type=request.report_type)
    display_title = _safe_title(request.title, "\u4f1a\u8bae\u8bed\u97f3\u7eaa\u8981")
    analysis_payload["title"] = display_title

    try:
        history_id = _save_audio_transcript_history(
            display_title,
            request.transcript_text,
            analysis_payload=analysis_payload,
        )
        _sync_action_items_for_history(
            history_id,
            title=display_title,
            date=str(analysis_payload.get("date", "") or ""),
            actions=analysis_payload.get("actions", []),
        )
    except Exception as exc:
        logger.exception("Failed to persist audio meeting history: title=%s", display_title)
        raise HTTPException(status_code=500, detail=f"\u4fdd\u5b58\u8bed\u97f3\u7eaa\u8981\u5386\u53f2\u5931\u8d25\uff1a{exc}") from exc

    audio_request = _build_audio_minutes_request(
        display_title,
        request.transcript_text,
        analysis_payload,
        report_type=request.report_type,
    )
    return AudioTranscriptResponse(
        title=display_title,
        file_name=display_title,
        transcript_text=request.transcript_text,
        markdown=_render_audio_minutes_markdown(audio_request),
        history_id=history_id,
        date=audio_request.date,
        weekly_period=audio_request.weekly_period,
        decisions=[item for item in audio_request.decisions],
        actions=[action for action in audio_request.actions],
        risks=[item for item in audio_request.risks],
        validation=analysis_payload.get("validation"),
        report_markdown=str(analysis_payload.get("report_markdown") or ""),
        report_type=request.report_type,
    )


@app.put("/api/history/{history_id}/audio-transcript", response_model=AudioTranscriptResponse)
def update_audio_transcript_record(history_id: int, request: AudioTranscriptDownloadRequest) -> AudioTranscriptResponse:
    record = get_meeting_history_record(history_id)
    if not record:
        raise HTTPException(status_code=404, detail="\u672a\u627e\u5230\u5bf9\u5e94\u7684\u5386\u53f2\u8bb0\u5f55")
    if str(record.get("entry_type") or "") != ENTRY_TYPE_AUDIO_TRANSCRIPT:
        raise HTTPException(status_code=400, detail="\u8be5\u5386\u53f2\u8bb0\u5f55\u4e0d\u662f\u8bed\u97f3\u7eaa\u8981\u8bb0\u5f55")

    analysis_payload = _run_workflow(request.transcript_text, report_type=request.report_type)
    display_title = _safe_title(request.title, "\u4f1a\u8bae\u8bed\u97f3\u7eaa\u8981")
    analysis_payload["title"] = display_title

    updated = _update_audio_transcript_history_entry(
        history_id,
        display_title,
        request.transcript_text,
        analysis_payload=analysis_payload,
    )
    if not updated:
        raise HTTPException(status_code=500, detail="\u66f4\u65b0\u8bed\u97f3\u7eaa\u8981\u5386\u53f2\u5931\u8d25")

    _sync_action_items_for_history(
        history_id,
        title=display_title,
        date=str(analysis_payload.get("date", "") or ""),
        actions=analysis_payload.get("actions", []),
    )

    audio_request = _build_audio_minutes_request(
        display_title,
        request.transcript_text,
        analysis_payload,
        report_type=request.report_type,
    )
    return AudioTranscriptResponse(
        title=display_title,
        file_name=display_title,
        transcript_text=request.transcript_text,
        markdown=_render_audio_minutes_markdown(audio_request),
        history_id=history_id,
        date=audio_request.date,
        weekly_period=audio_request.weekly_period,
        decisions=[item for item in audio_request.decisions],
        actions=[action for action in audio_request.actions],
        risks=[item for item in audio_request.risks],
        validation=analysis_payload.get("validation"),
        report_markdown=str(analysis_payload.get("report_markdown") or ""),
        report_type=request.report_type,
    )


@app.post("/api/report/weekly", response_model=WeeklyReportResponse)
def generate_weekly_report(request: WeeklyReportRequest) -> WeeklyReportResponse:
    state = _build_report_state(request)
    return WeeklyReportResponse(markdown=render_report_markdown(state))


@app.post("/api/report/weekly/markdown")
def download_weekly_report_markdown(request: WeeklyReportRequest) -> StreamingResponse:
    markdown = _render_text_meeting_markdown(request)
    filename = f"{_safe_title(request.title, '会议纪要')}_{request.report_type}.md"
    return _download_bytes_response(markdown.encode("utf-8"), "text/markdown; charset=utf-8", filename)


@app.post("/api/report/weekly/docx")
def download_weekly_report_docx(request: WeeklyReportRequest) -> StreamingResponse:
    filename = f"{_safe_title(request.title, '会议纪要')}_{request.report_type}.docx"
    return _download_bytes_response(
        _build_weekly_docx_bytes(request).getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename,
    )




@app.get("/api/action-items", response_model=list[ActionItemRecord])
def get_action_items() -> list[ActionItemRecord]:
    items = list_action_items()
    return [ActionItemRecord(**item) for item in items]


@app.patch("/api/action-items/{action_item_id}")
def update_action_item(action_item_id: int, request: ActionItemStatusUpdateRequest) -> dict[str, int | str]:
    try:
        updated = update_action_item_status(action_item_id, request.status)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    if not updated:
        raise HTTPException(status_code=404, detail="\u672a\u627e\u5230\u5bf9\u5e94\u7684\u884c\u52a8\u9879")
    return {"message": "\u66f4\u65b0\u6210\u529f", "updated_id": action_item_id, "status": request.status}


@app.delete("/api/action-items/{action_item_id}")
def remove_action_item(action_item_id: int) -> dict[str, int | str]:
    deleted = delete_action_item(action_item_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="\u672a\u627e\u5230\u5bf9\u5e94\u7684\u884c\u52a8\u9879")
    return {"message": "\u5220\u9664\u6210\u529f", "deleted_id": action_item_id}


@app.get("/api/history", response_model=list[MeetingHistoryItem])
def get_history() -> list[MeetingHistoryItem]:
    items = list_meeting_history()
    return [MeetingHistoryItem(**item) for item in items]


@app.delete("/api/history/{history_id}")
def remove_history_item(history_id: int) -> dict[str, int | str]:
    deleted = delete_meeting_history(history_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="未找到对应的历史记录")
    return {"message": "删除成功", "deleted_id": history_id}


@app.delete("/api/history")
def remove_all_history() -> dict[str, int | str]:
    deleted_count = clear_meeting_history()
    return {"message": "清空成功", "deleted_count": deleted_count}


@app.get("/api/history/{history_id}/download/{download_option}")
def download_history_file(history_id: int, download_option: str) -> StreamingResponse:
    record = get_meeting_history_record(history_id)
    if not record:
        raise HTTPException(status_code=404, detail="未找到对应的会议历史记录")

    option = download_option.strip().lower()
    entry_type = str(record.get("entry_type") or ENTRY_TYPE_MEETING_ANALYSIS)

    if entry_type == ENTRY_TYPE_AUDIO_TRANSCRIPT:
        return _download_audio_transcript_record(record, option)

    return _download_meeting_analysis_record(record, option)
